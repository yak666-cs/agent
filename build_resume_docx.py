from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT_PATH = r"C:\Users\DYK\Desktop\agent\agent-loop-lab\邓杨凯个人简历_AI Agent版.docx"


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("left", "top", "right", "bottom"):
        if edge in kwargs:
            edge_data = kwargs[edge]
            tag = f"w:{edge}"
            element = tc_borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tc_borders.append(element)
            for key in ("val", "sz", "space", "color"):
                if key in edge_data:
                    element.set(qn(f"w:{key}"), str(edge_data[key]))


def set_page(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.3)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)
    section.start_type = WD_SECTION.CONTINUOUS


def set_default_font(doc: Document):
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(10.5)


def add_run(paragraph, text, size=10.5, bold=False, color=None, name="Arial", east_asia="微软雅黑"):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def add_section_title(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    add_run(p, title, size=12, bold=True, color=(34, 34, 34))
    p_border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "9E9E9E")
    p_border.append(bottom)
    p._element.get_or_add_pPr().append(p_border)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    add_run(p, text, size=10.2)


def add_project(doc, title, subtitle, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    add_run(p, title, size=11, bold=True)
    if subtitle:
        add_run(p, f"  {subtitle}", size=10.2, color=(90, 90, 90))
    for bullet in bullets:
        add_bullet(doc, bullet)


doc = Document()
set_page(doc)
set_default_font(doc)

for style_name in ("List Bullet", "List Paragraph"):
    if style_name in doc.styles:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(10.2)

name = doc.add_paragraph()
name.paragraph_format.space_after = Pt(1)
name.alignment = WD_ALIGN_PARAGRAPH.LEFT
add_run(name, "邓杨凯", size=20, bold=True)

headline = doc.add_paragraph()
headline.paragraph_format.space_before = Pt(0)
headline.paragraph_format.space_after = Pt(2)
add_run(headline, "AI Agent 工程师 / Java 后端开发", size=11.5, color=(58, 58, 58))

contact = doc.add_paragraph()
contact.paragraph_format.space_before = Pt(0)
contact.paragraph_format.space_after = Pt(4)
add_run(contact, "13542021026 | 1853430587@qq.com | 广东广州 | 2003.05", size=10.2, color=(80, 80, 80))

summary = doc.add_paragraph()
summary.paragraph_format.space_before = Pt(0)
summary.paragraph_format.space_after = Pt(4)
summary.paragraph_format.line_spacing = 1.18
add_run(summary, "简介  ", size=10.5, bold=True)
add_run(
    summary,
    "面向 AI Agent 落地与后端工程开发，具备 Java 后端、工作流自动化、RAG 检索增强和工具调用型智能体项目经验。"
    "完成过慢病随访 AI 系统、桌面 Agent 与高并发购票系统开发，熟悉从需求拆解、接口设计、数据库建模、"
    "消息推送到部署运维的完整实现链路。",
    size=10.3,
)

add_section_title(doc, "教育背景")
edu = doc.add_paragraph()
edu.paragraph_format.space_before = Pt(2)
edu.paragraph_format.space_after = Pt(0)
add_run(edu, "华南理工大学", size=11, bold=True)
add_run(edu, "  2022.09 - 2026.06", size=10.2, color=(90, 90, 90))

edu2 = doc.add_paragraph()
edu2.paragraph_format.space_before = Pt(0)
edu2.paragraph_format.space_after = Pt(2)
edu2.paragraph_format.line_spacing = 1.15
add_run(edu2, "计算机科学与技术（本科）", size=10.4)
add_run(edu2, "  |  CET-6", size=10.2, color=(90, 90, 90))

course = doc.add_paragraph()
course.paragraph_format.space_before = Pt(0)
course.paragraph_format.space_after = Pt(4)
course.paragraph_format.line_spacing = 1.15
add_run(
    course,
    "主修课程：Java 程序设计、数据结构、算法设计与分析、计算机网络、操作系统、计算机组成与体系结构、深度学习与神经网络等。",
    size=10.2,
)

add_section_title(doc, "项目经历")
add_project(
    doc,
    "KAI Agent 桌面智能体",
    "独立项目",
    [
        "基于 DeepSeek V4-Flash、FastAPI 与 pywebview 构建桌面 AI Agent，支持多会话管理、文件上传解析、系统托盘运行与 Web/桌面双入口。",
        "设计 Agent 核心循环与工具体系，封装 bash、文件读写、文档读取、系统信息、进程管理、长期记忆等 10+ 内置工具，并实现 tool registry、权限分级与安全沙箱控制。",
        "实现长期记忆、对话持久化、上下文裁剪修复、Prompt Cache 分层缓存、重试熔断和子 Agent 分派，提升多轮对话稳定性与复杂任务执行能力。",
    ],
)
add_project(
    doc,
    "慢病患者智能随访与用药提醒系统",
    "毕业设计 / 全栈开发",
    [
        "基于 n8n + MySQL + 微信小程序设计并实现慢病随访闭环系统，覆盖患者端与医院端的用药查询、康复问卷、消息通知、出诊信息管理等核心场景。",
        "围绕论文方案完成患者认证、用药查询、定时用药提醒、动态问卷、健康资讯抓取、医院消息推送等工作流编排，并通过 Webhook 实现前后端数据交互。",
        "基于 n8n LangChain 节点、DeepSeek Chat、自定义 Tool、FAISS 向量库和 BGE-small-zh 嵌入模型搭建 AI 智能小助手与轻量级 RAG 知识库，提升医学问答的专业性与可用性。",
    ],
)
add_project(
    doc,
    "仿 12306 高并发购票系统",
    "购票模块开发",
    [
        "基于 Redis Lua 令牌桶实现库存原子扣减，在高并发购票场景下保障余票一致性与扣减准确性。",
        "结合本地锁与 Redisson 公平锁处理多实例并发竞争，避免同座位重复分配与超卖问题。",
        "设计订单超时/取消后的座位解锁、令牌回滚与余票恢复流程，降低占票不支付导致的资源阻塞风险。",
    ],
)

add_section_title(doc, "专业技能")
skills = [
    ("后端开发", "Java、Spring Boot、MyBatis、MySQL、Redis、JWT、PageHelper、RESTful API、Webhook 接口设计"),
    ("AI Agent / LLM", "n8n、LangChain、RAG、Tool Calling、Prompt Engineering、DeepSeek、Coze、FAISS、BGE-small-zh"),
    ("工程能力", "FastAPI、Python、SQLite、Docker、Git、Linux、定时任务、日志观测、缓存设计、异常重试与熔断"),
    ("客户端与业务", "微信小程序、患者/医院双端业务流程设计、问卷系统、消息推送、健康资讯抓取"),
]
for label, value in skills:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.12
    add_run(p, f"{label}：", size=10.2, bold=True)
    add_run(p, value, size=10.2)

add_section_title(doc, "自我评价")
self_eval = doc.add_paragraph()
self_eval.paragraph_format.space_before = Pt(2)
self_eval.paragraph_format.space_after = Pt(0)
self_eval.paragraph_format.line_spacing = 1.15
add_run(
    self_eval,
    "对 AI 应用和后端工程保持持续关注，具备较强的学习能力与责任心；能够快速理解业务需求并推进落地，"
    "同时具备良好的沟通协作意识，愿意在 AI Agent 工程化与 Java 后端方向持续深耕。",
    size=10.2,
)

doc.save(OUTPUT_PATH)
print(OUTPUT_PATH)
