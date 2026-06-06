"""
KAI AGENT —— 桌面应用

双击启动，自动运行 FastAPI 后端 + pywebview 原生窗口。
类 Cursor/Codex 桌面应用体验。
"""

import asyncio
import json
import os
import sys
import threading
import logging

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

# ── 提前加载 .env ──
env_path = os.path.join(os.path.dirname(__file__), ".env")
if os.path.exists(env_path):
    with open(env_path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                k, v = line.split("=", 1)
                k, v = k.strip(), v.strip().strip('"').strip("'")
                if k not in os.environ:
                    os.environ[k] = v

from agent.loop import AgentLoop
from agent.bootstrap import register_core_skills, register_core_tools
from agent.llm import LLMClient
from agent.context import ContextManager, DEFAULT_SYSTEM_PROMPT

from tools.registry import tool_registry
from tools.selector import ToolSelector
from tools.executor import ToolExecutor
from tools.builtin.bash import BashTool
from tools.builtin.file_reader import ReadFileTool
from tools.builtin.file_writer import WriteFileTool
from tools.builtin.web_search import WebSearchTool
from tools.builtin.python_repl import PythonReplTool
from tools.builtin.file_deleter import FileDeleterTool
from tools.builtin.system_info import SystemInfoTool
from tools.builtin.process_manager import ProcessManagerTool
from tools.builtin.read_document import ReadDocumentTool

from tools.builtin.ip_geolocation import IpGeolocationTool
from tools.builtin.memory_tool import MemoryTool

from skills.manager import skill_manager
from skills.builtin.system_debug import SystemDebugSkill
from skills.builtin.file_ops import FileOpsSkill
from skills.builtin.data_analysis import DataAnalysisSkill
from skills.builtin.file_manage import FileManageSkill
from skills.builtin.memory import MemorySkill

from harness.observability import AgentObserver, trace_store
from tools.prompt_cache import prompt_cache
from tools.sandbox import Sandbox, SandboxPolicy, SandboxMode
from harness.resilience import CircuitState
from harness.auth import AuthStore, get_current_user

logging.basicConfig(level=logging.WARNING, format="%(asctime)s [%(name)s] %(levelname)s: %(message)s")


def _register_all():
    register_core_tools()
    register_core_skills()


# 活跃会话管理
active_sessions: dict[str, AgentLoop] = {}

# 多会话持久化
from memory.conversation_store import ConversationStore, msg_to_row, row_to_msg
conversation_store = ConversationStore()

# ── 全局可配置沙箱（所有会话共享，前端可实时调整） ──
_sandbox_policy = SandboxPolicy(
    mode=SandboxMode.RESTRICTIVE,
    disabled_tools=["file_deleter"],
    allowed_commands=[
        "ls ", "dir ", "type ", "head ", "tail ",
        "find ", "grep ", "which ", "echo ", "pwd", "whoami",
        "df ", "du ", "ps ", "top ", "netstat ",
        "tasklist ", "systeminfo ", "ipconfig ", "ping ",
        "wmic ", "chdir ", "cd ", "tree ", "fc ", "comp ",
        "help ", "ver", "date ", "time ", "where ",
    ],
)
active_sandbox = Sandbox(_sandbox_policy)


def _sandbox_blocked_intent_reply(message: str) -> str | None:
    """Return an immediate user-facing sandbox denial for common blocked intents."""
    text = (message or "").lower()
    if "write_file" not in _sandbox_policy.disabled_tools:
        return None

    write_markers = [
        "写入", "创建", "新建", "保存", "生成", "建立",
        "write", "create", "save",
    ]
    file_markers = [
        "文件", "文档", ".txt", ".md", ".json", ".csv", ".py",
        "file", "document",
    ]

    if any(k in text for k in write_markers) and any(k in text for k in file_markers):
        return "沙盒当前已禁用文件写入工具 write_file，所以我不能创建或写入文件。"

    return None


def _build_sandbox_prompt() -> str:
    disabled = ", ".join(_sandbox_policy.disabled_tools) or "none"
    return (
        "\n\n## Runtime Sandbox Policy\n"
        f"Disabled tools: {disabled}.\n"
        "If a user request requires a disabled tool, do not claim the task is done. "
        "Explain that the sandbox has disabled that capability and name the disabled tool."
    )

# 上下文策略信息（每个 Agent 运行结束后更新）
_last_context_info: dict = {}
_session_trace_index: dict[str, str] = {}


def _store_context_info(ctx):
    """保存上下文策略信息供前端查询"""
    _last_context_info.clear()
    _last_context_info.update({
        "max_tokens": ctx.max_tokens,
        "reserve_tokens": ctx.reserve_tokens,
        "window_size": ctx.window_size,
        "enable_summary": ctx.enable_summary,
        "enable_relevance": ctx.enable_relevance,
        "compression_log": ctx.compression_log[-20:] if ctx.compression_log else [],
        "total_compressions": len(ctx.compression_log),
    })


def get_next_session_name() -> str:
    """自动分配不重复的名称：对话 1、对话 2……"""
    existing = conversation_store.list_sessions(999)
    used = set()
    for s in existing:
        name = s.get("name", "")
        if name.startswith("对话 ") and name[3:].isdigit():
            used.add(int(name[3:]))
    n = 1
    while n in used:
        n += 1
    return f"对话 {n}"


def gen_session_id() -> str:
    import time
    return f"s{int(time.time())}_{__import__('random').Random().randint(100000, 999999)}"


async def chat_stream(message: str, session_id: str = "", file_context: str = ""):
    queue = asyncio.Queue(maxsize=100)
    _register_all()

    def safe(data):
        return json.loads(json.dumps(data, ensure_ascii=False, default=str))

    blocked_reply = _sandbox_blocked_intent_reply(message)
    if blocked_reply:
        yield f"event: done\ndata: {json.dumps({'reply': blocked_reply}, ensure_ascii=False)}\n\n"
        return

    # 从数据库加载会话历史，转为 Message 对象
    history_msgs = []
    if session_id:
        from agent.types import Message as AgentMessage, Role
        try:
            stored = conversation_store.get_messages(session_id)
            for m in stored:
                msg = AgentMessage(
                    role=Role(m["role"]),
                    content=m["content"],
                )
                td = m.get("tool_data")
                if td:
                    if td.get("type") == "tool_calls":
                        msg.tool_calls = td.get("data", [])
                    elif td.get("type") == "tool_result":
                        msg.tool_call_id = td.get("tool_call_id", "")
                        msg.name = td.get("name", "")
                history_msgs.append(msg)
        except Exception:
            pass

        # 把用户消息立即落库，切回会话时能看到自己的提问
        try:
            conversation_store.save_messages(session_id, [
                {"role": "user", "content": message, "tool_data": None},
            ])
        except Exception:
            pass

    skill_prompt = skill_manager.build_skill_prompt() + _build_sandbox_prompt()
    llm = LLMClient()
    ctx = ContextManager(
        system_prompt=DEFAULT_SYSTEM_PROMPT,
        skill_prompt=skill_prompt,
        window_size=40,
        enable_summary=True,
        enable_relevance=True,
    )
    agent = AgentLoop(llm=llm, context=ctx, max_turns=12)
    agent.register_tool_registry(tool_registry)
    agent.register_tool_selector(ToolSelector(tool_registry))

    # 使用全局可配置沙箱（所有会话共享，前端可实时调整）
    agent.register_tool_executor(ToolExecutor(sandbox=active_sandbox))

    # Subagent 委派工具（LLM 自主决定何时分解复杂任务）
    from tools.builtin.subagent_tool import SubagentDelegateTool
    subagent_tool = SubagentDelegateTool(
        llm=llm,
        tool_registry=tool_registry,
        tool_executor=agent._tool_executor,
        tool_selector=agent._tool_selector,
        context_manager=ctx,
        event_cb=lambda evt, data: (
            observer.on_subagent_event(evt, data),
            queue.put_nowait(safe({"event": evt, "data": data})),
        )[-1],
    )
    tool_registry.register(subagent_tool)

    # 可观测性
    observer = AgentObserver(model=llm.model, session_id=session_id)
    if session_id:
        _session_trace_index[session_id] = observer.trace_id

    if session_id:
        active_sessions[session_id] = agent

    def on_turn_start(state):
        queue.put_nowait(safe({"event": "turn", "data": {"turn": state.turn_count}}))

    def on_status_change(state):
        queue.put_nowait(
            safe(
                {
                    "event": "agent_status",
                    "data": {
                        "turn": state.turn_count,
                        "status": state.status.value,
                        "message_count": len(state.messages),
                        "tool_calls": len(state.tool_call_history),
                        "tool_results": len(state.tool_result_history),
                    },
                }
            )
        )

    def on_tool_call(tc):
        queue.put_nowait(safe({"event": "tool_call", "data": {"name": tc.name, "arguments": tc.arguments}}))

    def on_tool_result(tr):
        queue.put_nowait(safe({
            "event": "tool_result",
            "data": {
                "name": tr.name,
                "status": "success" if tr.success else "failed",
                "output": (tr.output or tr.error or "")[:800],
            }
        }))

    agent.on_turn_start(lambda s: (on_turn_start(s), observer.on_turn_start(s)))
    agent.on_turn_end(observer.on_turn_end)
    agent.on_tool_call(lambda tc: (on_tool_call(tc), observer.on_tool_call(tc)))
    agent.on_tool_result(lambda tr: (on_tool_result(tr), observer.on_tool_result(tr)))
    agent.on_llm_response(observer.on_llm_response)
    agent.on_status_change(lambda s: (on_status_change(s), observer.on_status_change(s)))

    # ── 接下来的处理 ──
    augmented_message = f"{file_context}\n\n{message}" if file_context else message

    # 注入相关长期记忆
    try:
        from memory import MemoryStore
        mem_store = MemoryStore()
        mem_context = mem_store.get_relevant_context(message)
        if mem_context:
            augmented_message = f"{mem_context}\n\n---\n\n{augmented_message}"
    except Exception:
        pass

    async def run():
        try:
            result = await agent.run(augmented_message, previous_messages=history_msgs)

            # 更新上下文策略信息供前端查询
            _store_context_info(ctx)

            # 保存会话历史到数据库
            if session_id and hasattr(agent, '_last_state'):
                try:
                    rows = []
                    for m in agent._last_state.messages:
                        if m.role.name == "SYSTEM":
                            continue
                        tool_data = None
                        if m.tool_calls:
                            tool_data = {"type": "tool_calls", "data": m.tool_calls}
                        elif m.role.name == "TOOL":
                            tool_data = {"type": "tool_result",
                                         "tool_call_id": m.tool_call_id,
                                         "name": m.name}
                        rows.append({
                            "role": m.role.value,
                            "content": m.content,
                            "tool_data": tool_data,
                        })
                    if rows:
                        conversation_store.clear_session_messages(session_id)
                        conversation_store.save_messages(session_id, rows)
                except Exception:
                    pass

            metrics_dict = json.loads(observer.metrics.to_json()) if hasattr(observer, 'metrics') else {}
            cb = agent._circuit_breaker
            trace_payload = observer.export()
            queue.put_nowait(safe({"event": "done", "data": {
                "reply": result,
                "metrics": metrics_dict,
                "trace": trace_payload,
                "harness": {
                    "circuit_breaker": cb.state.value,
                    "cache_hit_rate": round(prompt_cache.stats.hit_rate, 3),
                    "cache_hits": prompt_cache.stats.hits,
                    "cache_misses": prompt_cache.stats.misses,
                    "sandbox_mode": agent._tool_executor.sandbox.policy.mode,
                },
            }}))
        except Exception as e:
            logging.exception("Agent 异常")
            queue.put_nowait(safe({"event": "error", "data": {"error": str(e)}}))

    asyncio.create_task(run())

    while True:
        try:
            event = await asyncio.wait_for(queue.get(), timeout=1.0)
            yield f"event: {event['event']}\ndata: {json.dumps(event['data'], ensure_ascii=False)}\n\n"
            if event['event'] in ('done', 'error'):
                break
        except asyncio.TimeoutError:
            yield ": keepalive\n\n"

    if session_id:
        active_sessions.pop(session_id, None)


# ── FastAPI ──
from fastapi import FastAPI, Request, UploadFile, File, Depends
from fastapi.responses import HTMLResponse, StreamingResponse, JSONResponse, FileResponse, Response
from fastapi.middleware.cors import CORSMiddleware
import pathlib

app = FastAPI(title="KAI AGENT")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])

UPLOAD_DIR = os.path.join(os.path.dirname(__file__), "uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)


@app.get("/icon.png")
async def icon():
    """桌面应用图标"""
    icon_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), "OIP-C.webp")
    if os.path.exists(icon_path):
        return FileResponse(icon_path, media_type="image/png")
    return Response(status_code=404)


def extract_document_text(path: str, ext: str) -> str:
    """上传时自动解析文档内容"""
    try:
        if ext == ".pdf":
            import fitz
            doc = fitz.open(path)
            total = len(doc)
            parts = [f"(PDF 文档, 共 {total} 页)"]
            for i in range(min(total, 50)):
                text = doc[i].get_text().strip()
                if text:
                    parts.append(f"--- 第 {i+1} 页 ---\n{text}")
            doc.close()
            text = "\n".join(parts)
            return text[:10000] + ("\n\n[内容过长已截断]" if len(text) > 10000 else "")

        elif ext == ".docx":
            from docx import Document
            doc = Document(path)
            lines = []
            for para in doc.paragraphs:
                if para.text.strip():
                    lines.append(para.text)
            text = "\n".join(lines)
            return text[:10000] + ("\n\n[内容过长已截断]" if len(text) > 10000 else "")

        elif ext in (".xlsx", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            parts = [f"(Excel 文档, {len(wb.sheetnames)} 个工作表)"]
            for name in wb.sheetnames[:5]:
                ws = wb[name]
                parts.append(f"--- 工作表: {name} ---")
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= 200:
                        parts.append("...(省略)")
                        break
                    row_vals = [str(v) if v is not None else "" for v in row]
                    line = "\t".join(row_vals)
                    if line.strip():
                        parts.append(line)
            wb.close()
            text = "\n".join(parts)
            return text[:10000] + ("\n\n[内容过长已截断]" if len(text) > 10000 else "")

    except ImportError as e:
        return f"(缺少依赖: {e})"
    except Exception as e:
        return f"(解析失败: {e})"
    return ""


SUPPORTED_EXTENSIONS = {
    # 文本与代码
    ".txt": "文本文档", ".md": "Markdown", ".py": "Python", ".js": "JavaScript",
    ".ts": "TypeScript", ".jsx": "React JSX", ".tsx": "React TSX",
    ".java": "Java", ".cpp": "C++", ".c": "C", ".h": "C 头文件", ".hpp": "C++ 头文件",
    ".go": "Go", ".rs": "Rust", ".rb": "Ruby", ".php": "PHP",
    ".swift": "Swift", ".kt": "Kotlin", ".scala": "Scala",
    ".sh": "Shell", ".bat": "批处理", ".ps1": "PowerShell",
    ".html": "HTML", ".css": "CSS", ".scss": "SCSS", ".less": "Less",
    ".json": "JSON", ".xml": "XML", ".yaml": "YAML", ".yml": "YAML",
    ".toml": "TOML", ".ini": "INI", ".cfg": "配置文件", ".conf": "配置文件",
    ".log": "日志", ".env": "环境变量",
    ".csv": "CSV", ".tsv": "TSV",
    ".sql": "SQL", ".r": "R", ".m": "MATLAB",
    # 文档
    ".pdf": "PDF 文档",
    ".docx": "Word 文档", ".doc": "Word 文档",
    ".xlsx": "Excel 表格", ".xls": "Excel 表格",
    # 图片（读取基本信息）
    ".png": "PNG 图片", ".jpg": "JPEG 图片", ".jpeg": "JPEG 图片",
    ".gif": "GIF 图片", ".webp": "WebP 图片", ".bmp": "BMP 图片", ".ico": "ICO 图标",
    ".svg": "SVG 矢量图",
}


@app.get("/api/supported-types")
async def get_supported_types():
    """返回支持的文件类型列表"""
    return JSONResponse(SUPPORTED_EXTENSIONS)


@app.post("/api/upload")
async def upload_file(file: UploadFile = File(...)):
    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        types_list = ", ".join(sorted(SUPPORTED_EXTENSIONS.keys()))
        return JSONResponse(
            {"error": f"不支持 {ext}，支持的类型有：{types_list}"},
            status_code=400,
        )

    import time
    safe_name = f"{int(time.time())}_{file.filename}"
    save_path = os.path.join(UPLOAD_DIR, safe_name)
    content = await file.read()
    with open(save_path, "wb") as f:
        f.write(content)

    target_path = os.path.normpath(save_path)

    # 文档类文件自动解析文本
    doc_text = ""
    doc_exts = {".pdf", ".docx", ".doc", ".xlsx", ".xls"}
    if ext in doc_exts:
        try:
            doc_text = extract_document_text(target_path, ext)
        except Exception as e:
            doc_text = f"(自动解析失败: {e})"

    resp = {
        "filename": file.filename,
        "saved_as": safe_name,
        "path": target_path,
        "size": len(content),
        "type": SUPPORTED_EXTENSIONS.get(ext, "未知"),
    }
    if doc_text:
        resp["content"] = doc_text
    return JSONResponse(resp)


# ── 多会话 API ──

@app.get("/api/sessions")
async def list_sessions():
    """列出所有会话（含消息数量）"""
    return conversation_store.list_sessions(50)


@app.get("/api/sessions/{sid}/messages")
async def get_session_messages(sid: str):
    """获取会话的消息列表"""
    msgs = conversation_store.get_messages(sid)
    return msgs


@app.post("/api/sessions")
async def create_session():
    """创建新会话"""
    sid = gen_session_id()
    name = get_next_session_name()
    conversation_store.create_session(sid, name)
    return {"id": sid, "name": name}


@app.delete("/api/sessions/{sid}")
async def delete_session(sid: str):
    conversation_store.delete_session(sid)
    return {"status": "ok"}


@app.put("/api/sessions/{sid}/rename")
async def rename_session(sid: str, name: str = ""):
    if name:
        ok = conversation_store.rename_session(sid, name)
        return {"status": "ok" if ok else "error"}
    return {"status": "error", "error": "名称为空"}


@app.get("/")
async def index():
    html_path = pathlib.Path(__file__).parent / "static" / "index.html"
    if html_path.exists():
        html = open(html_path, encoding="utf-8").read()
        html = html.replace("__ICON__", "/icon.png")
        return HTMLResponse(html)
    return HTMLResponse("<h1>KAI AGENT</h1><p>缺少 static/index.html</p>")


@app.get("/api/chat")
async def chat(request: Request, message: str = "", session_id: str = "", file_path: str = "", doc_content: str = ""):
    if not message:
        return StreamingResponse(
            (f"event: error\ndata: {json.dumps({'error': '消息不能为空'})}\n\n" for _ in [1]),
            media_type="text/event-stream",
        )

    # 如果 session 不存在，自动创建
    if session_id:
        existing = conversation_store.list_sessions(999)
        if not any(s["id"] == session_id for s in existing):
            conversation_store.create_session(session_id, get_next_session_name())

    # 如果有上传文件，在 Agent 上下文中注入文件信息
    file_context = ""
    if doc_content:
        # 文档内容已自动解析，直接注入
        fname = ""
        if file_path and os.path.exists(file_path):
            fname = os.path.basename(file_path)
        file_context = f"以下是用户上传的文档{f'（{fname}）' if fname else ''}内容：\n---\n{doc_content[:8000]}\n---"
    elif file_path and os.path.exists(file_path):
        fname = os.path.basename(file_path)
        fext = os.path.splitext(fname)[1].lower()
        fsize = os.path.getsize(file_path)
        ftype = SUPPORTED_EXTENSIONS.get(fext, "未知")
        size_str = f"{fsize/1024:.1f}KB" if fsize > 1024 else f"{fsize}B"
        doc_exts = {".pdf", ".docx", ".doc", ".xlsx", ".xls"}
        if fext in doc_exts:
            file_context = f"用户已上传文档：{fname}（{ftype}，{size_str}）\n文件绝对路径：{os.path.normpath(file_path)}\n请使用 read_document 工具读取此文件内容。"
        else:
            file_context = f"用户已上传文件：{fname}（{ftype}，{size_str}）\n文件绝对路径：{os.path.normpath(file_path)}\n如需处理该文件，请使用 read_file 工具读取其内容。"

    return StreamingResponse(
        chat_stream(message, session_id, file_context),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "Connection": "keep-alive"},
    )


@app.post("/api/cancel/{session_id}")
async def cancel_chat(session_id: str):
    agent = active_sessions.get(session_id)
    if agent:
        agent.cancel()
        return {"status": "ok", "message": "已发送中断请求"}
    return {"status": "not_found", "message": "未找到活跃会话"}


# ── 认证 API ──

auth_store = AuthStore()


@app.post("/api/register")
async def register(body: dict):
    """用户注册"""
    username = (body.get("username") or "").strip()
    password = body.get("password") or ""
    if not username or not password:
        return JSONResponse({"ok": False, "error": "用户名和密码不能为空"}, status_code=400)
    result = auth_store.register(username, password)
    if not result["ok"]:
        return JSONResponse(result, status_code=409)
    return result


@app.post("/api/login")
async def login(body: dict):
    """用户登录"""
    username = (body.get("username") or "").strip()
    password = body.get("password") or ""
    if not username or not password:
        return JSONResponse({"ok": False, "error": "用户名和密码不能为空"}, status_code=400)
    result = auth_store.login(username, password)
    if not result["ok"]:
        return JSONResponse(result, status_code=401)
    return result


@app.post("/api/logout")
async def logout(request: Request):
    """登出"""
    auth = request.headers.get("Authorization", "")
    if auth.startswith("Bearer "):
        auth_store.logout(auth[len("Bearer "):].strip())
    return {"ok": True}


@app.get("/api/context-info")
async def context_info():
    """获取当前上下文策略信息和压缩记录"""
    return _last_context_info if _last_context_info else {
        "max_tokens": 128000, "reserve_tokens": 4000,
        "window_size": 40, "enable_summary": True, "enable_relevance": True,
        "compression_log": [], "total_compressions": 0,
        "message": "还没有 Agent 运行记录，以上为默认配置"
    }


@app.get("/api/observability/{session_id}")
async def get_session_observability(session_id: str):
    trace_id = _session_trace_index.get(session_id, session_id)
    payload = trace_store.get(trace_id)
    if payload is None:
        return JSONResponse({"error": "trace_not_found", "session_id": session_id}, status_code=404)
    return JSONResponse(payload)


@app.get("/api/me")
async def me(user: dict = Depends(get_current_user)):
    """获取当前登录用户"""
    if user is None:
        return {"ok": False, "user": None}
    return {"ok": True, "user": {"id": user["id"], "username": user["username"]}}


# ── 沙箱配置 API ──


@app.get("/api/sandbox/config")
async def get_sandbox_config():
    """获取当前沙箱配置"""
    p = _sandbox_policy
    return {
        "mode": p.mode,
        "disabled_tools": p.disabled_tools,
        "allowed_commands": p.allowed_commands,
        "blocked_commands": p.blocked_commands,
        "write_blocked_files": p.write_blocked_files,
        "max_output_lines": p.max_output_lines,
        "max_output_chars": p.max_output_chars,
    }


@app.put("/api/sandbox/config")
async def update_sandbox_config(body: dict):
    """更新沙箱配置（仅更新请求中携带的字段）"""
    p = _sandbox_policy
    if "mode" in body and body["mode"] in (SandboxMode.RESTRICTIVE, SandboxMode.PERMISSIVE):
        p.mode = body["mode"]
    if "disabled_tools" in body and isinstance(body["disabled_tools"], list):
        p.disabled_tools = body["disabled_tools"]
    if "allowed_commands" in body and isinstance(body["allowed_commands"], list):
        p.allowed_commands = body["allowed_commands"]
    if "blocked_commands" in body and isinstance(body["blocked_commands"], list):
        p.blocked_commands = body["blocked_commands"]
    if "write_blocked_files" in body and isinstance(body["write_blocked_files"], list):
        p.write_blocked_files = body["write_blocked_files"]
    if "max_output_lines" in body and isinstance(body["max_output_lines"], int):
        p.max_output_lines = max(100, body["max_output_lines"])
    if "max_output_chars" in body and isinstance(body["max_output_chars"], int):
        p.max_output_chars = max(1000, body["max_output_chars"])
    return {"status": "ok", "mode": p.mode}


# ── 桌面窗口 ──
import socket

def _kill_old_instance(port: int):
    """杀掉占用指定端口的旧进程"""
    try:
        import subprocess
        result = subprocess.run(
            ["netstat", "-ano"], capture_output=True, text=True, timeout=5
        )
        for line in result.stdout.splitlines():
            if f"127.0.0.1:{port}" in line and "LISTENING" in line:
                parts = line.strip().split()
                pid = parts[-1]
                if pid.isdigit():
                    subprocess.run(["taskkill", "/F", "/PID", pid],
                                   capture_output=True, timeout=5)
    except Exception:
        pass


def start_server(port: int):
    import uvicorn

    class ReuseAddrServer(uvicorn.Server):
        def create_socket(self, family, type_, proto):
            sock = socket.socket(family, type_, proto)
            sock.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
            sock.bind((self.config.host, self.config.port))
            sock.setblocking(False)
            return sock

    # pythonw.exe 无控制台 sys.stdout=None，uvicorn 默认 log_config 会调用
    # sys.stdout.isatty() 导致 AttributeError。日志写入文件。
    log_file = os.path.join(os.path.dirname(__file__), "kai_server.log")
    config = uvicorn.Config(
        app, host="127.0.0.1", port=port,
        log_level="warning",
        log_config={
            "version": 1,
            "disable_existing_loggers": False,
            "formatters": {
                "plain": {"format": "%(asctime)s %(levelname)s %(message)s"},
            },
            "handlers": {
                "plain": {
                    "class": "logging.handlers.RotatingFileHandler",
                    "formatter": "plain",
                    "filename": log_file,
                    "maxBytes": 1048576,
                    "backupCount": 2,
                },
            },
            "loggers": {
                "uvicorn": {"handlers": ["plain"], "level": "WARNING"},
                "uvicorn.error": {"handlers": ["plain"], "level": "WARNING"},
            },
        },
    )
    server = ReuseAddrServer(config=config)
    server.run()


def main():
    port = int(os.environ.get("KAI_AGENT_PORT", 8765))

    _kill_old_instance(port)

    t = threading.Thread(target=start_server, args=(port,), daemon=True)
    t.start()

    # 轮询等待服务器就绪
    import time
    import urllib.request
    url = f"http://127.0.0.1:{port}"
    for _ in range(30):
        try:
            urllib.request.urlopen(url, timeout=1)
            break
        except Exception:
            time.sleep(0.5)

    # 读取前端页面，注入图标 URL 占位（具体图标由 FastAPI /icon.png 提供）
    html_path = os.path.join(os.path.dirname(__file__), "static", "index.html")
    with open(html_path, encoding="utf-8") as f:
        html = f.read()
    html = html.replace("__ICON__", "/icon.png")

    # pywebview 窗口 —— 直接加载服务器 URL，避免 about:blank 的跨域/CORS 问题
    import webview
    window = webview.create_window(
        "KAI AGENT",
        url=f"{url}/",
        width=900,
        height=700,
        resizable=True,
        min_size=(600, 400),
    )

    # ── 系统托盘 ──
    import pystray
    from PIL import Image

    ico_path = os.path.join(os.path.dirname(__file__), "kai_agent.ico")
    tray_img = Image.open(ico_path)

    def on_show(icon, item):
        """单击托盘图标：显示窗口"""
        try:
            window.show()
        except Exception:
            pass

    def on_quit(icon, item):
        """右键退出：彻底杀死进程"""
        icon.stop()
        os._exit(0)

    def _start_tray():
        icon = pystray.Icon(
            "kai_agent", tray_img, "KAI AGENT",
            menu=pystray.Menu(
                pystray.MenuItem("显示窗口", on_show, default=True),
                pystray.Menu.SEPARATOR,
                pystray.MenuItem("退出", on_quit),
            ),
        )
        icon.run()

    tray_thread = threading.Thread(target=_start_tray, daemon=False)
    tray_thread.start()

    # 点击 X 时隐藏到托盘而非退出
    def on_closing():
        try:
            window.hide()
        except Exception:
            pass
        return False  # 阻止窗口销毁，保留托盘

    window.events.closing += on_closing

    webview.start(gui="edgechromium")

    # webview.start 返回后等待托盘线程（用户点退出才会走到这里）
    tray_thread.join()


if __name__ == "__main__":
    main()
